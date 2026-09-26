#!/usr/bin/env python3
"""看一眼资料包里到底有什么。

用途：拿到新的景区资料包时，先跑这个了解结构，再决定怎么切块入库。
只用标准库 + openpyxl（docx 用 zipfile 直接解 XML，避免额外依赖）。
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path


def setup_console() -> None:
    if sys.platform == "win32":
        try:
            import ctypes

            ctypes.windll.kernel32.SetConsoleOutputCP(65001)
        except Exception:
            pass
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass


setup_console()


def _cell_text(cell_xml: str) -> str:
    parts = re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", cell_xml, flags=re.DOTALL)
    return "".join(parts).strip()


def docx_text(path: Path) -> tuple[list[str], list[list[list[str]]], int]:
    """从 docx 里抽段落与表格（直接解析 document.xml，不依赖 python-docx）。

    返回 (段落列表, 表格列表, 图片数)
    表格结构：tables[表序号][行序号][单元格序号] = 文本
    """
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")

    image_count = len([n for n in names if n.startswith("word/media/")])

    # 段落：<w:p>...</w:p> 里的所有 <w:t> 文本
    # 注意 \s 那一层：`<w:t>` 与 `<w:t xml:space="preserve">` 都匹配，
    # 但排除 `<w:tab/>` `<w:tbl>` 这类同前缀的自闭合/容器标签（否则会抽到 XML 片段）
    paragraphs: list[str] = []
    for para in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.DOTALL):
        texts = re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", para, flags=re.DOTALL)
        line = "".join(texts).strip()
        if line:
            paragraphs.append(line)

    # 表格：<w:tbl> → <w:tr> → <w:tc> → 单元格文本
    tables: list[list[list[str]]] = []
    for tbl in re.findall(r"<w:tbl>.*?</w:tbl>", xml, flags=re.DOTALL):
        rows: list[list[str]] = []
        for tr in re.findall(r"<w:tr[ >].*?</w:tr>", tbl, flags=re.DOTALL):
            cells = re.findall(r"<w:tc>.*?</w:tc>", tr, flags=re.DOTALL)
            rows.append([_cell_text(c) for c in cells])
        if rows:
            tables.append(rows)
    return paragraphs, tables, image_count


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: python inspect_corpus.py <资料包目录>")
        return 2

    root = Path(sys.argv[1])
    if not root.is_dir():
        print(f"目录不存在: {root}")
        return 2

    for f in sorted(root.rglob("*")):
        if not f.is_file():
            continue
        print("=" * 74)
        print(f"文件：{f.name}   ({f.stat().st_size / 1024:.1f} KB)")
        print("=" * 74)

        if f.suffix.lower() == ".docx":
            paras, tables, images = docx_text(f)
            print(f"  段落数：{len(paras)}   表格数：{len(tables)}   图片数：{images}")
            total = sum(len(p) for p in paras)
            print(f"  正文字数（估）：{total}")
            for ti, rows in enumerate(tables, start=1):
                cols = len(rows[0]) if rows else 0
                print(f"  ── 表{ti}：{len(rows)} 行 × {cols} 列  表头：{' | '.join(c[:10] for c in (rows[0] if rows else []))}")
                for ri, row in enumerate(rows[1:3], start=1):
                    print(f"     数据行{ri}: " + " | ".join(c[:20] for c in row[:5]))
            print("  ── 前 12 段 ──")
            for p in paras[:12]:
                print(f"    {p[:88]}")
            print("  ── 后 6 段 ──")
            for p in paras[-6:]:
                print(f"    {p[:88]}")

        elif f.suffix.lower() in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError:
                print("  (需要 openpyxl 才能看 xlsx，跳过)")
                continue
            wb = load_workbook(f, read_only=True, data_only=True)
            print(f"  工作表：{wb.sheetnames}")
            for name in wb.sheetnames:
                ws = wb[name]
                print(f"    [{name}] {ws.max_row} 行 × {ws.max_column} 列")
                for i, row in enumerate(ws.iter_rows(max_row=3, values_only=True)):
                    cells = [("" if c is None else str(c))[:16] for c in row[:10]]
                    print(f"      行{i+1}: {cells}")
            wb.close()

        else:
            print("  (未处理的类型)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
